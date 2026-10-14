import io
from docx import Document
from fastapi.testclient import TestClient


def create_dummy_docx() -> bytes:
    doc = Document()
    doc.add_heading("Paritosh Dash", 0)
    doc.add_paragraph("Email: paritosh@example.com | Phone: +91 9876543210")
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph("Full-stack engineer with expertise in Python, FastAPI, Docker, and PostgreSQL.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI, PostgreSQL, Docker, React, Machine Learning, Redis")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.Tech in Computer Science from National Institute of Technology")

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_resume_upload_and_versioning(client: TestClient, auth_headers: dict):
    docx_bytes = create_dummy_docx()

    # Upload v1
    files = {"file": ("resume_v1.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res1 = client.post("/api/v1/resumes/upload", files=files, headers=auth_headers)
    assert res1.status_code == 201
    v1_data = res1.json()
    assert v1_data["is_current"] is True
    assert v1_data["has_parsed_text"] is True
    assert "Python" in v1_data["extracted_skills"]
    assert "FastAPI" in v1_data["extracted_skills"]
    v1_id = v1_data["id"]

    # Upload v2 (version replacement)
    files2 = {"file": ("resume_v2.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res2 = client.post("/api/v1/resumes/upload", files=files2, headers=auth_headers)
    assert res2.status_code == 201
    v2_data = res2.json()
    assert v2_data["version"] > v1_data["version"]
    assert v2_data["is_current"] is True

    # Check list of resumes
    list_res = client.get("/api/v1/resumes", headers=auth_headers)
    assert list_res.status_code == 200
    res_list = list_res.json()
    assert len(res_list) >= 2

    # Check current resume endpoint
    cur_res = client.get("/api/v1/resumes/current", headers=auth_headers)
    assert cur_res.status_code == 200
    assert cur_res.json()["id"] == v2_data["id"]

    # Test download endpoint
    dl_res = client.get(f"/api/v1/resumes/{v1_id}/download", headers=auth_headers)
    assert dl_res.status_code == 200
    assert len(dl_res.content) > 0
